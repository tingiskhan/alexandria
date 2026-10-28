"""Real Docling on the formats that need no models (no OCR, no layout), so these run without a GPU.

The files are built here with the libraries Docling itself reads them with, instead of being checked in.
"""

from io import BytesIO

import docx
import openpyxl
import pptx
import pytest

from alexandria_core.models import Page
from alexandria_ingest.extract import DoclingExtractor


@pytest.fixture(scope="module")
def extractor() -> DoclingExtractor:
    return DoclingExtractor({"docx", "pptx", "xlsx", "html", "md", "txt"})


def test_markdown_is_one_page(extractor: DoclingExtractor) -> None:
    pages = extractor.extract(b"# Plan\n\nShip it by Friday.", "plan.md")

    assert pages == [Page(1, "# Plan\n\nShip it by Friday.")]


def test_text_is_read_as_markdown(extractor: DoclingExtractor) -> None:
    pages = extractor.extract(b"Plain notes.\n\nS&P rose.", "notes.txt")

    assert pages == [Page(1, "Plain notes.\n\nS&P rose.")]


def test_html_is_one_page_without_escaping(extractor: DoclingExtractor) -> None:
    pages = extractor.extract(b"<html><body><h1>Report</h1><p>S&amp;P rose.</p></body></html>", "report.html")

    assert pages == [Page(1, "# Report\n\nS&P rose.")]


def test_word_is_one_page_with_headings(extractor: DoclingExtractor) -> None:
    document = docx.Document()
    document.add_heading("Minutes", 1)
    document.add_paragraph("Budget approved.")
    content = BytesIO()
    document.save(content)

    [page] = extractor.extract(content.getvalue(), "minutes.docx")

    assert page.number == 1
    assert "Minutes" in page.markdown
    assert "Budget approved." in page.markdown


def test_powerpoint_has_a_page_per_slide(extractor: DoclingExtractor) -> None:
    presentation = pptx.Presentation()
    for title in ("First slide", "Second slide"):
        presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = title
    content = BytesIO()
    presentation.save(content)

    pages = extractor.extract(content.getvalue(), "deck.pptx")

    assert [p.number for p in pages] == [1, 2]
    assert "First slide" in pages[0].markdown
    assert "Second slide" in pages[1].markdown


def test_excel_has_a_page_per_sheet_with_tables(extractor: DoclingExtractor) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.append(["region", "revenue"])
    workbook.active.append(["north", 4471])
    workbook.create_sheet("Notes").append(["draft"])
    content = BytesIO()
    workbook.save(content)

    pages = extractor.extract(content.getvalue(), "figures.xlsx")

    assert [p.number for p in pages] == [1, 2]
    assert "| region" in pages[0].markdown
    assert "4471" in pages[0].markdown
    assert "draft" in pages[1].markdown
